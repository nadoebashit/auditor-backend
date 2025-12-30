from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, status
from sqlalchemy.orm import Session
from docx import Document

from app.core.config import settings
from app.core.db import get_db
from app.core.helpers.helpers import get_arq_redis
from app.modules.auth.models import User
from app.modules.auth.router import get_current_user
from app.modules.files.models import FileIndexStatus, FileScope, StoredFile
from app.modules.files.storage import FileStorage, S3Config
from app.modules.projects.models import (
    GeneratedDocument,
    LegalMatterRecord,
    MaterialityRecord,
    PBCItem,
    Project,
    RiskRecord,
)
from app.modules.actions.schemas import (
    ActionResponse,
    AddLegalMatterRequest,
    AddRiskRequest,
    GenerateDocumentRequest,
    SaveMaterialityRequest,
    UpdatePBCRequest,
)
from app.modules.projects.service import ensure_project_access


router = APIRouter(prefix="/actions", tags=["actions"])


def _get_storage() -> FileStorage:
    return FileStorage(
        S3Config(
            endpoint_url=settings.S3_ENDPOINT_URL,
            access_key=settings.S3_ACCESS_KEY,
            secret_key=settings.S3_SECRET_KEY,
            region=settings.S3_REGION,
            bucket_admin_laws=settings.S3_BUCKET_ADMIN_LAWS,
            bucket_customer_docs=settings.S3_BUCKET_CUSTOMER_DOCS,
        )
    )


@router.post("/save-materiality", response_model=ActionResponse)
def save_materiality(
    payload: SaveMaterialityRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = ensure_project_access(db, payload.project_id, current_user)

    # Upsert: one per project
    existing = (
        db.query(MaterialityRecord)
        .filter(MaterialityRecord.project_id == project.id)
        .first()
    )

    if existing is None:
        rec = MaterialityRecord(project_id=project.id)
        db.add(rec)
    else:
        rec = existing

    rec.benchmark = payload.benchmark
    rec.benchmark_value = payload.benchmark_value
    rec.risk_level = payload.risk_level
    rec.overall_materiality = payload.om
    rec.performance_materiality = payload.pm
    rec.clearly_trivial_threshold = payload.ct
    rec.rationale = payload.rationale

    rec.confirmed_by = current_user.id
    rec.confirmed_at = datetime.utcnow()

    db.commit()
    db.refresh(rec)

    return ActionResponse(success=True, message="Существенность сохранена", record_id=rec.id)


@router.post("/add-risk", response_model=ActionResponse)
def add_risk(
    payload: AddRiskRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = ensure_project_access(db, payload.project_id, current_user)

    rec = RiskRecord(
        project_id=project.id,
        cycle=payload.cycle,
        assertion=payload.assertion,
        risk_description=payload.risk_description,
        inherent_risk=payload.inherent_risk,
        control_risk=payload.control_risk,
        detection_risk=payload.detection_risk,
        response=payload.response,
        is_significant=bool(payload.is_significant),
        is_fraud_risk=bool(payload.is_fraud_risk),
        confirmed_by=current_user.id,
        confirmed_at=datetime.utcnow(),
    )

    db.add(rec)
    db.commit()
    db.refresh(rec)

    return ActionResponse(success=True, message="Риск добавлен", record_id=rec.id)


@router.post("/add-legal-matter", response_model=ActionResponse)
def add_legal_matter(
    payload: AddLegalMatterRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = ensure_project_access(db, payload.project_id, current_user)

    rec = LegalMatterRecord(
        project_id=project.id,
        matter_name=payload.matter_name,
        claim_amount=payload.claim_amount,
        probability=payload.probability,
        outcome_estimable=bool(payload.outcome_estimable),
        is_material=bool(payload.is_material),
        disclosure_required=bool(payload.disclosure_required),
        provision_required=bool(payload.provision_required),
        is_kam=bool(payload.is_kam),
        rationale=payload.rationale,
        confirmed_by=current_user.id,
        confirmed_at=datetime.utcnow(),
    )

    db.add(rec)
    db.commit()
    db.refresh(rec)

    return ActionResponse(success=True, message="Юридический вопрос сохранён", record_id=rec.id)


@router.post("/update-pbc", response_model=ActionResponse)
def update_pbc(
    payload: UpdatePBCRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = ensure_project_access(db, payload.project_id, current_user)

    item = (
        db.query(PBCItem)
        .filter(PBCItem.project_id == project.id, PBCItem.item_code == payload.item_code)
        .first()
    )

    if item is None:
        item = PBCItem(project_id=project.id, item_code=payload.item_code, item_name=payload.item_name or payload.item_code)
        db.add(item)

    if payload.item_name is not None:
        item.item_name = payload.item_name
    if payload.cycle is not None:
        item.cycle = payload.cycle
    if payload.priority is not None:
        item.priority = payload.priority
    if payload.status is not None:
        item.status = payload.status
    if payload.due_date is not None:
        item.due_date = payload.due_date
    if payload.received_date is not None:
        item.received_date = payload.received_date
    if payload.notes is not None:
        item.notes = payload.notes

    db.commit()
    db.refresh(item)

    return ActionResponse(success=True, message="PBC обновлён", record_id=item.id)


@router.post("/generate-document", response_model=ActionResponse)
async def generate_document(
    payload: GenerateDocumentRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project: Project = ensure_project_access(db, payload.project_id, current_user)

    tenant_id = str(project.id)

    # Minimal docx generation (MVP): render key-value pairs.
    doc = Document()
    doc.add_heading(f"Generated document {payload.template_id}", level=1)
    doc.add_paragraph(f"Project: {project.project_code} | Client: {project.client_name}")
    doc.add_paragraph(f"Generated at: {datetime.utcnow().isoformat()}Z")
    doc.add_paragraph(" ")

    data = payload.data or {}
    if isinstance(data, dict) and data:
        table = doc.add_table(rows=1, cols=2)
        hdr = table.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Value"
        for k, v in data.items():
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(v)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    storage = _get_storage()

    filename = f"{payload.template_id}_{project.project_code}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.docx"
    object_key = f"outputs/{filename}"

    storage.upload_customer_file(
        customer_id=tenant_id,
        file_obj=buf,
        object_key=object_key,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    stored_file = StoredFile(
        owner_id=current_user.id,
        customer_id=tenant_id,
        scope=FileScope.CUSTOMER_DOC,
        bucket=settings.S3_BUCKET_CUSTOMER_DOCS,
        object_key=f"{tenant_id}/{object_key}",
        original_filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=int(buf.getbuffer().nbytes),
        is_indexed=False,
        index_error=None,
        index_status=FileIndexStatus.QUEUED,
    )
    db.add(stored_file)
    db.commit()
    db.refresh(stored_file)

    gen = GeneratedDocument(
        project_id=project.id,
        template_id=payload.template_id,
        output_format=payload.output_format,
        stored_file_id=stored_file.id,
        filename=filename,
        file_path=stored_file.object_key,
        mime_type=stored_file.content_type,
        size_bytes=stored_file.size_bytes,
        generated_by=current_user.id,
    )
    db.add(gen)
    db.commit()
    db.refresh(gen)

    # Best-effort: enqueue indexing of generated output (optional)
    try:
        redis = await get_arq_redis()
        await redis.enqueue_job("index_file_task", str(stored_file.id))
    except Exception:
        pass

    return ActionResponse(
        success=True,
        message="Документ сгенерирован",
        record_id=gen.id,
        extra={
            "stored_file_id": str(stored_file.id),
            "download_url": f"/api/v1/files/{stored_file.id}/download",
        },
    )
